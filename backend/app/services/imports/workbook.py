"""Build the synthetic workbook from the canonical records."""

from __future__ import annotations

from decimal import Decimal
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook

from app.domain.synthetic import SYNTHETIC_IDEAS
from app.services.savings import combined_annual_saving

VOLUME_COLUMNS = [
    ("HCV Rigid Truck (Annual Volume)", "HCV_RIGID"),
    ("HCV Tipper (Annual Volume)", "HCV_TIPPER"),
    ("Tractor-Trailer (Annual Volume)", "TRACTOR_TRAILER"),
    ("RMC (Annual Volume)", "RMC"),
    ("COWL (Annual Volume)", "COWL"),
]

HEADERS = [
    "Idea_ID",
    "Component",
    "Idea generation date",
    "Idea belonged department",
    "Rejection date",
    "Rejected department",
    *[name for name, _code in VOLUME_COLUMNS],
    "Idea_Description",
    "Current design or process",
    "Proposed technical change",
    "Engineering rationale",
    "Expected benefit",
    "Constraints or validation requirements",
    "Stage_Status",
    "Rejection_Category",
    "Rejection_Reason",
    "Idea_Type",
    "Applicable Vehicle / Tonnage Details",
    "Saving per vehicle (INR)",
    "Total Combined Annual Saving Per Idea (INR)",
    "Idea Disposition",
]


def idea_row(record: dict) -> list:
    volumes = []
    pairs = []
    for _name, code in VOLUME_COLUMNS:
        amount = Decimal(str(record["volumes"].get(code, 0)))
        volumes.append(float(amount))
        pairs.append((code, amount))
    total, _warnings = combined_annual_saving(Decimal(str(record["saving_per_vehicle"])), pairs)
    reported = record.get("reported_total_saving")
    reported_value = float(reported) if reported is not None else float(total or 0)
    return [
        record["idea_code"],
        record["component"],
        record["idea_generation_date"],
        record["idea_belonged_department"],
        record["rejection_date"],
        record["rejected_department"],
        *volumes,
        record["description"],
        record["current_design"],
        record["proposed_change"],
        record["engineering_rationale"],
        record["expected_benefit"],
        record["constraints"],
        record["stage_status"],
        record["rejection_category"],
        record["rejection_reason"],
        record["idea_type"],
        record["applicable_vehicle_details"],
        record["saving_per_vehicle"],
        reported_value,
        record["idea_disposition"],
    ]


def exception_rows() -> list[list]:
    base = idea_row(SYNTHETIC_IDEAS[0])
    missing_id = list(base)
    missing_id[0] = None
    duplicate_a = list(base)
    duplicate_a[0] = "VAVE-EX-DUP"
    duplicate_b = list(base)
    duplicate_b[0] = "VAVE-EX-DUP"
    duplicate_b[1] = "Different component"
    bad_stage = list(base)
    bad_stage[0] = "VAVE-EX-STAGE"
    bad_stage[HEADERS.index("Stage_Status")] = "Duplicate"
    bad_date = list(base)
    bad_date[0] = "VAVE-EX-DATE"
    bad_date[HEADERS.index("Idea generation date")] = "32-13-2024"
    bad_saving = list(base)
    bad_saving[0] = "VAVE-EX-SAVE"
    bad_saving[HEADERS.index("Saving per vehicle (INR)")] = "abc"
    missing_description = list(base)
    missing_description[0] = "VAVE-EX-DESC"
    missing_description[HEADERS.index("Idea_Description")] = None
    missing_department = list(base)
    missing_department[0] = "VAVE-EX-DEPT"
    missing_department[HEADERS.index("Idea belonged department")] = None
    reason_without_category = list(base)
    reason_without_category[0] = "VAVE-EX-REASON"
    reason_without_category[HEADERS.index("Rejection_Category")] = None
    reason_without_category[HEADERS.index("Rejection_Reason")] = "Duplicate"
    reason_without_category[HEADERS.index("Rejected department")] = "Quality"
    negative_volume = list(base)
    negative_volume[0] = "VAVE-EX-VOL"
    negative_volume[HEADERS.index("HCV Rigid Truck (Annual Volume)")] = -5
    return [
        missing_id,
        duplicate_a,
        duplicate_b,
        bad_stage,
        bad_date,
        bad_saving,
        missing_description,
        missing_department,
        reason_without_category,
        negative_volume,
    ]


def build_workbook_bytes() -> bytes:
    book = Workbook()
    cover = book.active
    cover.title = "Cover"
    cover["A1"] = "SYNTHETIC DATA"
    cover["A2"] = "These rows are fictional VAVE examples for the VAVE Intelligence prototype."
    cover["A3"] = "They are not company records and must not be described as production history."
    ideas = book.create_sheet("Ideas")
    ideas.append(HEADERS)
    for record in SYNTHETIC_IDEAS:
        ideas.append(idea_row(record))
    exceptions = book.create_sheet("Exceptions")
    exceptions.append(HEADERS)
    for row in exception_rows():
        exceptions.append(row)
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def write_workbook(path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(build_workbook_bytes())
    return path
