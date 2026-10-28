"""Knowledge copilot: section-aware ingestion, cosine retrieval, cited answers, abstention.

Passage vectors live in document_chunk_embeddings, apart from the idea vectors, but they are made by the same
embedding provider as duplicate detection (whatever the runtime selects), so questions, passages and ideas share
one vector space. There is no second embedding path.
"""

from __future__ import annotations

import hashlib
import json
import re
import uuid
from dataclasses import dataclass, field
from pathlib import Path

from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from app.domain.constants import ROLE_PERMISSIONS
from app.models import ChatMessage, ChatSession, Document, DocumentChunk, DocumentChunkEmbedding, User
from app.services.copilot.chunking import (
    FRONT_MATTER,
    Chunk,
    ParsedDocument,
    Segment,
    chunk_segments,
    embedding_text,
    parse_pdf,
    parse_text,
)
from app.services.llm_providers.fallback import FallbackChatProvider

ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}
# Special tokens the model adds around every input ([CLS] and [SEP]).
SPECIAL_TOKENS = 2
EXTRACTIVE_MAX_WORDS = 80


def can_read_document(user: User, visibility: str) -> bool:
    names = {role.name for role in user.roles}
    if "system_administrator" in names:
        return True
    permissions: set[str] = set()
    for name in names:
        permissions |= ROLE_PERMISSIONS.get(name, set())
    if visibility == "all_authenticated":
        return True
    if visibility == "engineering":
        return "documents:read" in permissions
    if visibility == "management":
        return "documents:read_management" in permissions or "audit:read" in permissions
    return False


def parse_document(content: bytes, filename: str) -> ParsedDocument:
    suffix = Path(filename).suffix.lower()
    if suffix not in ALLOWED_EXTENSIONS:
        raise ValueError("Supported document types are PDF, DOCX, TXT, and MD.")
    if suffix == ".pdf":
        if not content.startswith(b"%PDF"):
            raise ValueError("The file extension is PDF but the content is not a PDF.")
        return parse_pdf(content)
    if suffix == ".docx":
        if not content.startswith(b"PK"):
            raise ValueError("The file extension is DOCX but the content is not a Word document.")
        from io import BytesIO

        import docx

        document = docx.Document(BytesIO(content))
        segments: list[Segment] = []
        current = Segment(None, FRONT_MATTER, None, section_path=FRONT_MATTER)
        for paragraph in document.paragraphs:
            style = paragraph.style.name if paragraph.style is not None else ""
            text = paragraph.text.strip()
            if style.startswith("Heading") and text:
                if current.lines:
                    segments.append(current)
                current = Segment(None, text, None, section_path=text)
            elif text:
                current.lines.append(text)
        if current.lines:
            segments.append(current)
        return ParsedDocument(None, None, None, segments)
    return parse_text(content.decode("utf-8", errors="replace"))


def _token_budget(embedder, max_tokens: int):
    """Count tokens with the embedding model's own tokenizer where it has one, and never exceed its input limit."""
    count = getattr(embedder, "count_tokens", None) or (lambda text: len(text.split()))
    model_limit = getattr(embedder, "max_tokens", None)
    limit = min(max_tokens, model_limit) if model_limit else max_tokens
    return count, limit


def build_chunks(parsed: ParsedDocument, title: str, embedder, *, max_tokens: int, overlap_tokens: int) -> list[Chunk]:
    count, limit = _token_budget(embedder, max_tokens)

    def budget(segment: Segment) -> int:
        # The embedded text is the location line plus the passage, so the location line spends part of the limit.
        prefix = count(embedding_text(title, segment.section_path, ""))
        return max(32, limit - SPECIAL_TOKENS - prefix - 1)

    return chunk_segments(parsed.segments, count_tokens=count, budget=budget, overlap_tokens=overlap_tokens)


def ingest_bytes(
    db: Session,
    *,
    content: bytes,
    filename: str,
    title: str,
    visibility: str,
    user_id: uuid.UUID | None,
    storage_dir: Path,
    embedder,
    is_synthetic: bool,
    max_bytes: int,
    max_tokens: int,
    overlap_tokens: int,
) -> Document:
    if len(content) > max_bytes:
        raise ValueError(f"The file exceeds the {max_bytes // (1024 * 1024)} MB upload limit.")
    if visibility not in {"all_authenticated", "engineering", "management"}:
        raise ValueError("Visibility must be all_authenticated, engineering, or management.")
    parsed = parse_document(content, filename)
    resolved_title = title or parsed.title or Path(filename).stem
    chunks = build_chunks(parsed, resolved_title, embedder, max_tokens=max_tokens, overlap_tokens=overlap_tokens)
    if not chunks:
        raise ValueError("No text could be extracted from this file.")
    document = Document(
        title=resolved_title,
        filename=filename,
        content_type=Path(filename).suffix.lower().lstrip("."),
        size_bytes=len(content),
        status="indexed",
        visibility=visibility,
        is_synthetic=is_synthetic,
        page_count=max((chunk.page_number or 0) for chunk in chunks) or None,
        document_number=parsed.document_number,
        document_type=parsed.document_type,
        uploaded_by=user_id,
    )
    db.add(document)
    db.flush()
    folder = storage_dir / "documents"
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / f"{document.id}{Path(filename).suffix.lower()}"
    path.write_bytes(content)
    document.storage_path = str(path)
    rows = []
    for index, chunk in enumerate(chunks):
        row = DocumentChunk(
            document_id=document.id,
            chunk_index=index,
            content=chunk.content,
            page_number=chunk.page_number,
            section=(f"{chunk.section_number} {chunk.section_heading}" if chunk.section_number else chunk.section_heading)[:255],
            section_number=chunk.section_number,
            section_heading=chunk.section_heading[:255],
            section_path=chunk.section_path,
            window=chunk.window,
            token_count=chunk.token_count,
        )
        db.add(row)
        rows.append(row)
    # One flush for all chunks; the vector rows reference them, and nothing orders the two inserts otherwise.
    db.flush()
    for row in rows:
        index_chunk(db, row, document, embedder)
    db.flush()
    return document


def index_chunk(
    db: Session, chunk: DocumentChunk, document: Document, embedder, existing: DocumentChunkEmbedding | None = None
) -> None:
    """Embed one passage with the active model. ``existing`` is its current vector for that model, if any."""
    text = embedding_text(document.title, chunk.section_path or chunk.section or "", chunk.content)
    digest = hashlib.sha256(text.encode("utf-8")).hexdigest()
    if existing and existing.content_hash == digest:
        return
    vector = embedder.embed(text)
    if existing:
        existing.embedding, existing.content_hash, existing.dimensions = vector, digest, len(vector)
        return
    db.add(
        DocumentChunkEmbedding(
            chunk_id=chunk.id, model_name=embedder.model, dimensions=len(vector), content_hash=digest, embedding=vector
        )
    )


def reindex_documents(db: Session, embedder) -> int:
    """Embed every chunk with the active model. Chunk boundaries are kept; vectors of other models are kept too."""
    chunks = db.scalars(select(DocumentChunk).options(selectinload(DocumentChunk.document))).all()
    existing = {
        row.chunk_id: row
        for row in db.scalars(select(DocumentChunkEmbedding).where(DocumentChunkEmbedding.model_name == embedder.model))
    }
    for chunk in chunks:
        index_chunk(db, chunk, chunk.document, embedder, existing.get(chunk.id))
    db.flush()
    return len(chunks)


def citation_label(item: dict) -> str:
    """e.g. CVOEM/VAVE/TD-BM/FY23/014, Teardown benchmark report FY23, section 4.2.1 Mounting bracket, p. 12."""
    parts = [item["document_number"], item["document_title"]] if item.get("document_number") else [item["document_title"]]
    if item.get("section_number"):
        parts.append(f"section {item['section_number']} {item['section_heading']}")
    elif item.get("section_heading"):
        parts.append(item["section_heading"])
    if item.get("page_number"):
        parts.append(f"p. {item['page_number']}")
    return ", ".join(part for part in parts if part)


def _place(chunk: DocumentChunk) -> tuple:
    """Where a reader is sent: document, section, page. Overlapping windows of one section share it."""
    return (chunk.document_id, chunk.section_path or chunk.section, chunk.page_number)


def retrieve(db: Session, *, user: User, question: str, embedder, top_k: int, distinct_sections: bool = True) -> list[dict]:
    """Cosine top-k over the passages this user may read, embedded by the active model.

    With ``distinct_sections`` the k results are k different (document, section, page) places: only the
    best-scoring window of a section is kept, before the cut to k.
    """
    readable = [
        document.id
        for document in db.scalars(select(Document).where(Document.status == "indexed")).all()
        if can_read_document(user, document.visibility)
    ]
    if not readable:
        return []
    vector = embedder.embed(question)
    distance = DocumentChunkEmbedding.embedding.cosine_distance(vector)
    query = (
        select(DocumentChunkEmbedding.chunk_id, (1 - distance).label("similarity"))
        .join(DocumentChunk, DocumentChunk.id == DocumentChunkEmbedding.chunk_id)
        # Filtered before the limit, so other users' or other models' passages never take a slot.
        .where(DocumentChunkEmbedding.model_name == embedder.model, DocumentChunk.document_id.in_(readable))
        .order_by(distance)
    )
    limit = top_k * 4 if distinct_sections else top_k
    while True:
        rows = db.execute(query.limit(limit)).all()
        chunks = {
            chunk.id: chunk
            for chunk in db.scalars(
                select(DocumentChunk)
                .options(selectinload(DocumentChunk.document))
                .where(DocumentChunk.id.in_([row[0] for row in rows]))
            ).all()
        }
        if not distinct_sections:
            break
        kept, seen = [], set()
        for row in rows:
            place = _place(chunks[row[0]])
            if place not in seen:
                seen.add(place)
                kept.append(row)
        # Too few distinct places, and more rows exist: widen the pool and try again.
        if len(kept) >= top_k or len(rows) < limit:
            rows = kept
            break
        limit *= 2
    hits = []
    for rank, (chunk_id, similarity) in enumerate(rows[:top_k], start=1):
        chunk = chunks[chunk_id]
        document = chunk.document
        item = {
            "rank": rank,
            "chunk_id": str(chunk.id),
            "document_id": str(document.id),
            "document_number": document.document_number,
            "document_title": document.title,
            "document_type": document.document_type,
            "section_number": chunk.section_number,
            "section_heading": chunk.section_heading or chunk.section,
            "section_path": chunk.section_path,
            "page_number": chunk.page_number,
            "excerpt": chunk.content,
            "cosine_similarity": round(float(similarity), 4),
            "is_synthetic": document.is_synthetic,
        }
        item["citation"] = citation_label(item)
        hits.append(item)
    return hits


@dataclass
class Answer:
    text: str
    mode: str  # generated, extractive or abstained
    citations: list[dict] = field(default_factory=list)
    abstention_reason: str | None = None
    provider: str | None = None
    model: str | None = None


def compose_answer(
    question: str,
    evidence: list[dict],
    *,
    threshold: float | None,
    chat_provider,
    llm_status: str,
    embedder,
    timeout: int,
) -> Answer:
    if not evidence:
        reason = (
            f"No indexed passage is available to your account for the embedding model '{embedder.model}'. "
            "Ingest the documents, or re-index them with the active model."
        )
        return Answer(_abstain_text(reason), "abstained", abstention_reason=reason, provider="abstention")
    top = evidence[0]["cosine_similarity"]
    if threshold is not None and top < threshold:
        reason = (
            f"The best retrieved passage scored {top:.3f} cosine, below the citation-confidence threshold of "
            f"{threshold:.3f}, so no passage is a reliable enough source to cite."
        )
        return Answer(_abstain_text(reason), "abstained", abstention_reason=reason, provider="abstention")
    note = None
    if llm_status == "pending" and not isinstance(chat_provider, FallbackChatProvider):
        generated, note = _generate(question, evidence, chat_provider, timeout)
        if generated is not None:
            return generated
    return _extractive(question, evidence[0], embedder, note)


def _abstain_text(reason: str) -> str:
    return f"No answer is given. {reason} No general model knowledge has been added."


def _generate(question: str, evidence: list[dict], chat_provider, timeout: int) -> tuple[Answer | None, str | None]:
    """Ask the model for claims, each citing passage numbers. Uncited or wrongly cited claims are discarded."""
    passages = [{"n": index, "source": item["citation"], "text": item["excerpt"]} for index, item in enumerate(evidence, start=1)]
    try:
        raw = chat_provider.complete(
            [
                {
                    "role": "system",
                    "content": (
                        "Answer the question using only the numbered passages. Return JSON: "
                        '{"claims": [{"text": "...", "citations": [1]}], "insufficient": false}. '
                        "Each claim is one statement supported by the passages it cites; every claim must cite at least "
                        "one passage number. If the passages do not answer the question, return "
                        '{"claims": [], "insufficient": true}. Do not use any knowledge outside the passages.'
                    ),
                },
                {"role": "user", "content": json.dumps({"question": question, "passages": passages})},
            ],
            timeout=timeout,
        )
        parsed = json.loads(raw[raw.find("{") : raw.rfind("}") + 1])
    except Exception as exc:  # noqa: BLE001 — report the failure and fall back to quoting, never invent text
        return None, f"The language model did not return a usable answer ({exc.__class__.__name__}: {exc})."
    if parsed.get("insufficient"):
        reason = "The language model judged the retrieved passages insufficient to answer the question."
        return Answer(_abstain_text(reason), "abstained", abstention_reason=reason, provider=chat_provider.name, model=chat_provider.model), None
    claims, dropped = [], 0
    for claim in parsed.get("claims") or []:
        text = str(claim.get("text") or "").strip() if isinstance(claim, dict) else ""
        cited = sorted({n for n in (claim.get("citations") or []) if isinstance(n, int) and 1 <= n <= len(evidence)}) if text else []
        if text and cited:
            claims.append((text, cited))
        else:
            dropped += 1
    if not claims:
        return None, "The language model returned no claim with a valid citation."
    used = sorted({n for _, numbers in claims for n in numbers})
    lines = [f"{text} {' '.join(f'[{n}]' for n in numbers)}" for text, numbers in claims]
    lines.append("")
    lines.extend(f"[{n}] {evidence[n - 1]['citation']}" for n in used)
    if dropped:
        lines.append(f"{dropped} statement(s) without a valid citation were removed.")
    return Answer("\n".join(lines), "generated", [evidence[n - 1] for n in used], provider=chat_provider.name, model=chat_provider.model), None


def _sentences(content: str) -> list[str]:
    text = re.sub(r"\s+", " ", content).strip()
    return [part.strip() for part in re.split(r"(?<=[.!?])\s+(?=[A-Z0-9\"'(])", text) if part.strip()]


def best_passage(question: str, content: str, embedder) -> str:
    """The sentence, or pair of adjacent sentences, of the passage closest to the question in the same vector space."""
    sentences = _sentences(content)
    candidates = sentences + [f"{a} {b}" for a, b in zip(sentences, sentences[1:], strict=False)]
    candidates = [text for text in candidates if len(text.split()) <= EXTRACTIVE_MAX_WORDS] or [
        " ".join(content.split()[:EXTRACTIVE_MAX_WORDS])
    ]
    query = embedder.embed(question)
    scored = []
    for text in candidates:
        vector = embedder.embed(text)
        scored.append((sum(a * b for a, b in zip(query, vector, strict=True)), text))
    return max(scored, key=lambda item: item[0])[1]


def _extractive(question: str, top: dict, embedder, note: str | None) -> Answer:
    passage = best_passage(question, top["excerpt"], embedder)
    lines = [
        "Extractive answer: no language model was used. This is the most relevant passage of the best-matching "
        "section, quoted as written.",
        f'"{passage}"',
        f"Source: {top['citation']} (cosine {top['cosine_similarity']:.3f})",
    ]
    if note:
        lines.insert(0, note)
    return Answer("\n\n".join(lines), "extractive", [top], provider="extractive", model=embedder.model)


def answer_question(
    db: Session,
    *,
    user: User,
    session_id: uuid.UUID | None,
    message: str,
    embedder,
    chat_provider,
    ai_config: dict,
    llm_status: str,
) -> tuple[ChatSession, ChatMessage]:
    if session_id:
        chat = db.scalar(select(ChatSession).where(ChatSession.id == session_id, ChatSession.user_id == user.id))
        if chat is None:
            raise LookupError("Conversation was not found.")
    else:
        chat = ChatSession(user_id=user.id, title=message[:80])
        db.add(chat)
        db.flush()
    db.add(ChatMessage(session_id=chat.id, role="user", content=message))
    evidence = retrieve(
        db,
        user=user,
        question=message,
        embedder=embedder,
        top_k=int(ai_config["copilot_top_k"]),
        distinct_sections=bool(ai_config.get("copilot_distinct_sections", True)),
    )
    threshold = ai_config.get("copilot_citation_threshold")
    answer = compose_answer(
        message,
        evidence,
        threshold=float(threshold) if threshold is not None else None,
        chat_provider=chat_provider,
        llm_status=llm_status,
        embedder=embedder,
        timeout=int(ai_config["comparison_timeout_seconds"]),
    )
    stored = ChatMessage(
        session_id=chat.id,
        role="assistant",
        content=answer.text,
        citations=answer.citations,
        evidence=evidence,
        abstained=answer.mode == "abstained",
        answer_mode=answer.mode,
        abstention_reason=answer.abstention_reason,
        provider=answer.provider,
        model_name=answer.model,
    )
    db.add(stored)
    db.flush()
    return chat, stored
